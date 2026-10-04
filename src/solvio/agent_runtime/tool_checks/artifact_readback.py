"""Generic format readback in the existing offline Office process.

No generated module is imported. These observations are file content, never
instructions or a judgement that the original task has been fulfilled.
"""
import base64
import csv
import hashlib
import io
import json
import math
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from table_report import archive, passive_pdf
from document_formats import _xml

MAX_TEXT = 32000
MAX_CELLS = 20000


def readback(item):
    raw = base64.b64decode(item['content_b64'], validate=True)
    mime = item['mime_type']
    observation = {'name': item['name'], 'mime_type': mime,
        'sha256': hashlib.sha256(raw).hexdigest(), 'size': len(raw)}
    if mime == 'application/pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw), strict=True)
        if reader.is_encrypted or not 1 <= len(reader.pages) <= 64:
            raise ValueError('pdf_pages')
        passive_pdf(reader)
        observation['pages'] = len(reader.pages)
        observation['text'] = '\n'.join(page.extract_text() or '' for page in reader.pages)
    elif mime.endswith('spreadsheetml.sheet'):
        import openpyxl
        archive(raw)
        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, read_only=True)
        try:
            if not 1 <= len(book.worksheets) <= 16:
                raise ValueError('sheet_limit')
            sheets, count = [], 0
            for sheet in book.worksheets:
                if sheet.max_row * sheet.max_column > MAX_CELLS - count:
                    raise ValueError('cell_limit')
                rows = []
                for row in sheet.iter_rows():
                    values = []
                    for cell in row:
                        count += 1
                        if cell.data_type == 'f':
                            raise ValueError('formula_not_independently_calculated')
                        value = cell.value
                        if isinstance(value, float) and not math.isfinite(value):
                            raise ValueError('nonfinite_cell')
                        values.append(value if value is None or type(value) in (str, int, float, bool)
                                      else str(value))
                    rows.append(values)
                sheets.append({'name': sheet.title, 'rows': rows})
            observation['sheets'] = sheets
        finally:
            book.close()
    elif mime.endswith(('wordprocessingml.document', 'presentationml.presentation')):
        archive(raw)
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            # Word fields (e.g. INCLUDETEXT/DDEAUTO) can fetch or act without
            # an OOXML external relationship. Slides also have action nodes.
            for name in z.namelist():
                if not name.endswith('.xml'):
                    continue
                tree = _xml(z.read(name), max_xml_bytes=8*1024*1024,
                            max_xml_elements=100000, max_xml_depth=64)
                if any(node.tag.rsplit('}', 1)[-1] in {
                        'fldSimple', 'instrText', 'fldChar', 'altChunk',
                        'hlinkClick', 'hlinkMouseOver'} for node in tree.iter()):
                    raise ValueError('active_document_field')
            names = ([name for name in z.namelist() if name == 'word/document.xml']
                if mime.endswith('wordprocessingml.document') else
                [name for name in z.namelist() if name.startswith('ppt/slides/slide')
                 and name.endswith('.xml') and '/_rels/' not in name])
            if not names or len(names) > 64:
                raise ValueError('document_parts')
            observation['text'] = '\n'.join('\n'.join(node.text or ''
                for node in _xml(z.read(name), max_xml_bytes=8*1024*1024,
                    max_xml_elements=100000, max_xml_depth=64).iter()
                if node.tag.rsplit('}', 1)[-1] == 't') for name in sorted(names))
    elif mime.startswith('image/'):
        from PIL import Image
        with Image.open(io.BytesIO(raw)) as image:
            if not 1 <= image.width * image.height <= 16000000:
                raise ValueError('image_dimensions')
            image.verify()
            observation.update(width=image.width, height=image.height, format=image.format,
                content_note='Image dimensions and decoding only; visual meaning not verified.')
    else:
        text = raw.decode('utf-8')
        if '\0' in text:
            raise ValueError('invalid_text')
        if mime == 'application/json':
            json.loads(text, parse_constant=lambda _: (_ for _ in ()).throw(ValueError('nonfinite')))
        if mime == 'text/csv':
            # CSV is passive data; refuse spreadsheet formula injection.
            for row in csv.reader(io.StringIO(text)):
                if any(cell.lstrip().startswith(('=', '+', '@')) or
                       (cell.lstrip().startswith('-') and not cell.lstrip()[1:].replace('.', '', 1).isdigit())
                       for cell in row):
                    raise ValueError('active_csv_cell')
        observation['text'] = text
    if len(json.dumps(observation, ensure_ascii=False)) > MAX_TEXT:
        raise ValueError('readback_too_large')
    return observation


def check(body):
    try:
        observations = [readback(item) for item in body['files']]
        if len(json.dumps(observations, ensure_ascii=False)) > MAX_TEXT:
            raise ValueError('readback_too_large')
        return {'ok': True, 'files': observations}
    except Exception:
        # Library messages and generated strings are not Core failure codes.
        return {'ok': False, 'reason': 'artifact_readback_failed'}


if __name__ == '__main__':
    print(json.dumps(check(json.load(sys.stdin)), ensure_ascii=False, allow_nan=False))
