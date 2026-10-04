"""Core-owned readback check, run ONLY in the isolated Office Python runtime.

Existing pandas/openpyxl/Pillow/pypdf parse actual inputs and outputs. No
generated module is imported. This is an output oracle, not an Office writer.
"""
import base64
import csv
import hashlib
import io
import json
import math
import posixpath
import re
import stat
import sys
import zipfile
from pathlib import Path
from urllib.parse import unquote, urlsplit

# Only this explicitly hashed stdlib module accompanies the Core-owned oracle
# into its private snapshot. No package/Core path or generated file is loaded.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from document_formats import _read_entry, _xml

MAX_ROWS = 50000
MAX_COLUMNS = 128
MAX_CELLS = 500000
MAX_REPORT_TEXT = 10000


# Only these Core-owned semantic failures leave the checker. Library exceptions
# (including ValueError with coincidentally matching text) are never forwarded.
class CheckRefused(ValueError):
    pass


REASONS = frozenset(('active_report', 'active_result_cell', 'active_workbook', 'archive_entry', 'archive_limit', 'archive_stream_invalid', 'cell_limit', 'chart_blank', 'chart_dimensions', 'csv_dialect', 'csv_header', 'empty_sheet', 'encoding', 'external_workbook_relationship', 'nonfinite_cell', 'report_object_limit', 'report_pages', 'report_text_limit', 'report_statistics_missing', 'report_table_missing', 'report_title_missing', 'result_sheet_names', 'sheet_limit', 'source_formulas_require_recalculation', 'statistics_differ_from_source', 'table_limit', 'table_shape', 'unbound_workbook_relationship', 'unsupported_table_format', 'workbook_chart_missing', 'workbook_source_data_changed', 'workbook_statistics_changed', 'workbook_xml_invalid'))
STAGES = frozenset({"input", "source", "statistics", "workbook", "chart", "report"})


def data(value):
    raw = base64.b64decode(value, validate=True)
    if base64.b64encode(raw).decode("ascii") != value:
        raise CheckRefused("encoding")
    return raw


def archive(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        entries = z.infolist()
        if (not 1 <= len(entries) <= 2048 or sum(e.file_size for e in entries) > 64 * 1024 * 1024
                or any(e.flag_bits & 1 or e.filename.startswith("/")
                       or ".." in e.filename.split("/") for e in entries)):
            raise CheckRefused("archive_limit")
        active_parts = ("vbaproject", "externallinks/", "embeddings/", "connections.xml",
                        "querytables/", "activex/", "ctrlprops/", "customui/")
        if any(any(part in e.filename.lower() for part in active_parts) for e in entries):
            raise CheckRefused("active_workbook")
        all_names = {entry.filename for entry in entries}
        names, offsets, expanded = set(), set(), 0
        for entry in entries:
            mode = stat.S_IFMT(entry.external_attr >> 16)
            if (entry.filename.casefold() in names or entry.header_offset in offsets
                    or entry.filename != entry.orig_filename or "\\" in entry.filename
                    or any(ord(ch) < 32 for ch in entry.filename)
                    or any(part in ("", ".", "..") for part in entry.filename.rstrip("/").split("/"))
                    or mode not in (0, stat.S_IFREG, stat.S_IFDIR)
                    or entry.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED)
                    or entry.flag_bits & 0x41 or entry.file_size < 0
                    or entry.compress_size < 0 or (entry.is_dir() and entry.file_size)):
                raise CheckRefused("archive_entry")
            names.add(entry.filename.casefold())
            offsets.add(entry.header_offset)
            try:
                content = _read_entry(z, entry, raw, 64 * 1024 * 1024 - expanded,
                    max_expanded_bytes=64 * 1024 * 1024, max_xml_bytes=64 * 1024 * 1024)
            except ValueError:
                raise CheckRefused("archive_stream_invalid") from None
            expanded += len(content)
            if entry.filename.lower().endswith((".xml", ".rels")):
                try:
                    tree = _xml(content, max_xml_bytes=64 * 1024 * 1024,
                                max_xml_elements=2_100_000, max_xml_depth=64)
                except ValueError:
                    raise CheckRefused("workbook_xml_invalid") from None
                for node in tree.iter():
                    local = node.tag.rsplit("}", 1)[-1]
                    if (local in {"connection", "connections", "queryTable", "queryTables", "oleObject",
                            "oleObjects", "control", "controls", "externalLink", "externalBook", "ddeLink",
                            "externalData", "dbPr", "webPr"}
                            or "macroenabled" in node.get("ContentType", "").lower()):
                        raise CheckRefused("active_workbook")
                    if local == "Relationship":
                        relation = node.get("Type", "").lower()
                        if (node.get("TargetMode", "Internal") != "Internal"
                                or relation.endswith(("/externallink", "/connections", "/querytable",
                                                      "/oleobject", "/control", "/vbaproject"))):
                            raise CheckRefused("external_workbook_relationship")
                        target = unquote(node.get("Target", ""))
                        parsed = urlsplit(target)
                        if (not target or parsed.scheme or parsed.netloc or parsed.query or "\\" in target
                                or any(ord(character) < 32 for character in target)):
                            raise CheckRefused("external_workbook_relationship")
                        base = ("" if entry.filename == "_rels/.rels"
                                else entry.filename.rsplit("/_rels/", 1)[0])
                        resolved = posixpath.normpath(parsed.path.lstrip("/") if parsed.path.startswith("/")
                            else posixpath.join(base, parsed.path))
                        if resolved not in all_names or resolved.startswith("../"):
                            raise CheckRefused("unbound_workbook_relationship")


def clean(value):
    import pandas as pd
    if value is None or pd.isna(value):
        return None
    if type(value) is bool:
        return value
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise CheckRefused("nonfinite_cell")
        return value
    if hasattr(value, "item"):
        return clean(value.item())
    if hasattr(value, "isoformat"):
        return value.isoformat()
    value = str(value)
    if len(value) > 32767:
        raise CheckRefused("cell_limit")
    return value


def sources(files):
    import pandas as pd
    import openpyxl
    result, cells = [], 0
    for item in files:
        raw = data(item["content_b64"])
        name = item["name"]
        if name.lower().endswith(".csv"):
            # pandas repairs duplicate/empty headers before exposing columns.
            # Validate the actual CSV header first so repair is not admission.
            text = raw.decode("utf-8-sig")
            try:
                separator = csv.Sniffer().sniff(text[:65536], delimiters=",;").delimiter
            except csv.Error:
                if "," in text[:65536] or ";" in text[:65536]:
                    raise CheckRefused("csv_dialect") from None
                separator = ","  # An ordinary single-column CSV has no separator.
            header = next(csv.reader(io.StringIO(text, newline=""), delimiter=separator, strict=True), [])
            if (not header or len(header) > MAX_COLUMNS or len(set(header)) != len(header)
                    or any(not value.strip() or len(value) > 120 for value in header)):
                raise CheckRefused("csv_header")
            # The existing parser receives the same measured dialect.
            frame = pd.read_csv(io.BytesIO(raw), encoding="utf-8-sig", sep=separator,
                engine="python", nrows=MAX_ROWS + 1)
            tables = [("CSV", frame)]
        elif name.lower().endswith(".xlsx"):
            archive(raw)
            book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            tables = []
            if not 1 <= len(book.worksheets) <= 8:
                raise CheckRefused("sheet_limit")
            for sheet in book.worksheets:
                sheet.reset_dimensions()
                values = []
                worksheet_cells = sum(frame.size for _, frame in tables)
                for index, row in enumerate(sheet.iter_rows()):
                    if index > MAX_ROWS or len(row) > MAX_COLUMNS:
                        raise CheckRefused("table_limit")
                    if index:
                        worksheet_cells += len(row)
                        if cells + worksheet_cells > MAX_CELLS:
                            raise CheckRefused("table_limit")
                    if any(cell.data_type == "f" for cell in row):
                        raise CheckRefused("source_formulas_require_recalculation")
                    values.append([cell.value for cell in row])
                if not values:
                    raise CheckRefused("empty_sheet")
                tables.append((sheet.title, pd.DataFrame(values[1:], columns=values[0])))
            book.close()
        else:
            raise CheckRefused("unsupported_table_format")
        for title, frame in tables:
            columns = list(frame.columns)
            if (not 1 <= len(columns) <= MAX_COLUMNS or len(frame) > MAX_ROWS
                    or any(type(c) is not str or not c.strip() or len(c) > 120 for c in columns)
                    or len(set(columns)) != len(columns)):
                raise CheckRefused("table_shape")
            cells += len(frame) * len(columns)
            if cells > MAX_CELLS or len(result) >= 8:
                raise CheckRefused("table_limit")
            numbers = {}
            for column in columns:
                series = frame[column]
                if pd.api.types.is_numeric_dtype(series) and not pd.api.types.is_bool_dtype(series):
                    usable = series.dropna()
                    if len(usable):
                        values = [float(v) for v in usable]
                        if not all(math.isfinite(v) for v in values):
                            raise CheckRefused("nonfinite_cell")
                        total = math.fsum(values)
                        numbers[column] = {"count": len(values), "sum": total,
                            "min": min(values), "max": max(values), "mean": total / len(values)}
            fact = {"source": name, "sheet": title, "rows": len(frame), "columns": columns,
                "missing": {c: int(frame[c].isna().sum()) for c in columns}, "numeric": numbers}
            rows = [[clean(value) for value in row]
                    for row in frame.itertuples(index=False, name=None)]
            result.append((fact, rows))
    return result


def passive_pdf(reader):
    """Inspect pypdf's resolved object graph, without inventing a PDF parser."""
    from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject, NameObject, StreamObject
    forbidden_keys = {"/A", "/AA", "/OpenAction", "/JS", "/JavaScript", "/AcroForm",
        "/XFA", "/EmbeddedFiles", "/EF", "/AF", "/Collection", "/RichMediaContent",
        "/RichMediaSettings", "/Sound", "/Movie", "/Rendition", "/Launch", "/URI"}
    forbidden_names = {"/JavaScript", "/Launch", "/URI", "/GoToR", "/GoToE", "/SubmitForm",
        "/ImportData", "/Rendition", "/Movie", "/Sound", "/RichMedia", "/FileAttachment",
        "/EmbeddedFile", "/Filespec", "/Screen", "/3D", "/Widget", "/Action"}
    pending, seen, count = [(reader.trailer, 0)], set(), 0
    while pending:
        node, depth = pending.pop()
        if depth > 64:
            raise CheckRefused("report_object_limit")
        if isinstance(node, IndirectObject):
            key = ("ref", node.idnum, node.generation)
            if key in seen:
                continue
            seen.add(key)
            node = node.get_object()
        if isinstance(node, (DictionaryObject, ArrayObject)):
            key = ("object", id(node))
            if key in seen:
                continue
            seen.add(key)
        count += 1
        if count > 50000:
            raise CheckRefused("report_object_limit")
        if isinstance(node, DictionaryObject):
            if forbidden_keys.intersection(node) or (isinstance(node, StreamObject) and "/F" in node):
                raise CheckRefused("active_report")
            pending.extend((value, depth + 1) for value in node.values())
        elif isinstance(node, ArrayObject):
            pending.extend((value, depth + 1) for value in node)
        elif isinstance(node, NameObject) and str(node) in forbidden_names:
            raise CheckRefused("active_report")


def equal(actual, expected):
    if type(expected) is float:
        return type(actual) in (int, float) and math.isfinite(actual) and math.isclose(
            actual, expected, rel_tol=1e-9, abs_tol=1e-9)
    if type(expected) is dict:
        return type(actual) is dict and set(actual) == set(expected) and all(
            equal(actual[k], v) for k, v in expected.items())
    if type(expected) is list:
        return type(actual) is list and len(actual) == len(expected) and all(
            equal(a, e) for a, e in zip(actual, expected))
    return type(actual) is type(expected) and actual == expected


def _check(payload, progress):
    import openpyxl
    from PIL import Image
    from pypdf import PdfReader
    progress[0] = "source"
    measured = sources(payload["input"]["files"])
    expected = [fact for fact, _ in measured]
    result = payload["result"]
    progress[0] = "statistics"
    if not equal(result["tables"], expected):
        raise CheckRefused("statistics_differ_from_source")
    progress[0] = "workbook"
    outputs = {item["name"]: data(item["content_b64"]) for item in result["files"]}
    raw = outputs["Analyse.xlsx"]
    archive(raw)
    book = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, read_only=False)
    expected_names = ["Kennzahlen"] + [f"Daten{i}" for i in range(1, len(measured) + 1)]
    if book.sheetnames != expected_names:
        raise CheckRefused("result_sheet_names")
    stats = [["Datei", "Tabelle", "Spalte", "Anzahl", "Summe", "Minimum", "Maximum", "Mittelwert"]]
    for index, (fact, rows) in enumerate(measured, 1):
        sheet = book[f"Daten{index}"]
        wanted = [fact["columns"]] + rows
        actual = [[cell.value for cell in row] for row in sheet.iter_rows()]
        if not equal(actual, wanted):
            raise CheckRefused("workbook_source_data_changed")
        for column, n in fact["numeric"].items():
            stats.append([fact["source"], fact["sheet"], column,
                n["count"], n["sum"], n["min"], n["max"], n["mean"]])
    stat_sheet = book["Kennzahlen"]
    actual = [[cell.value for cell in row] for row in stat_sheet.iter_rows()]
    if not equal(actual, stats):
        raise CheckRefused("workbook_statistics_changed")
    embedded_chart_count = sum(len(sheet._charts) for sheet in book.worksheets)
    if embedded_chart_count < 1:
        raise CheckRefused("workbook_chart_missing")
    if any(cell.data_type == "f" or cell.hyperlink for sheet in book.worksheets
           for row in sheet.iter_rows() for cell in row):
        raise CheckRefused("active_result_cell")
    book.close()
    progress[0] = "chart"
    image = Image.open(io.BytesIO(outputs["Diagramm.png"]))
    if image.format != "PNG" or not 320 <= image.width <= 2400 or not 200 <= image.height <= 1600:
        raise CheckRefused("chart_dimensions")
    image.load()
    if all(low == high for low, high in image.convert("RGB").getextrema()):
        raise CheckRefused("chart_blank")
    progress[0] = "report"
    reader = PdfReader(io.BytesIO(outputs["Bericht.pdf"]), strict=True)
    if reader.is_encrypted or not 1 <= len(reader.pages) <= 64:
        raise CheckRefused("report_pages")
    passive_pdf(reader)
    texts, text_length = [], 0
    for page in reader.pages:
        part = page.extract_text() or ""
        text_length += len(part) + (1 if texts else 0)
        if text_length > MAX_REPORT_TEXT:
            raise CheckRefused("report_text_limit")
        texts.append(part)
    text = "\n".join(texts)
    if "SOLVIO Tabellenbericht" not in text:
        raise CheckRefused("report_title_missing")
    sections = list(re.finditer(r"(?m)^Tabelle ([1-9][0-9]*): ([0-9]+) Zeilen\s*$", text))
    if len(sections) != len(expected):
        raise CheckRefused("report_table_missing")
    for index, (fact, section) in enumerate(zip(expected, sections), 1):
        if section.group(1) != str(index) or section.group(2) != str(fact["rows"]):
            raise CheckRefused("report_table_missing")
        end = sections[index].start() if index < len(sections) else len(text)
        table_text = text[section.end():end]
        for column, values in fact["numeric"].items():
            # A visible original column position is useful to readers and
            # remains verifiable when a font cannot display a Unicode label.
            marker = f"(Spalte {fact['columns'].index(column)+1}): Summe: {values['sum']:.6g}"
            if marker not in table_text:
                raise CheckRefused("report_statistics_missing")
    # These properties are measured above from the actual bytes. They do not
    # assert chart-data correctness, readability, or completion of a user goal.
    readback = {"version": 1,
        "workbook": {"statistics_match": True, "source_rows_match": True,
                     "embedded_chart_count": embedded_chart_count},
        "chart": {"format": image.format, "width": image.width,
                  "height": image.height, "nonblank": True},
        "report": {"sha256": hashlib.sha256(outputs["Bericht.pdf"]).hexdigest(),
                   "pages": len(reader.pages), "text": text,
                   "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest()}}
    return {"ok": True, "tables": expected, "readback": readback}


def check(payload):
    progress = ["input"]
    try:
        return _check(payload, progress)
    except CheckRefused as exc:
        reason = exc.args[0] if len(exc.args) == 1 and exc.args[0] in REASONS else "library_parse_failed"
    except Exception:
        reason = "library_parse_failed"
    return {"ok": False, "reason": reason, "stage": progress[0]}


if __name__ == "__main__":
    try:
        body = json.load(sys.stdin)
    except Exception:
        answer = {"ok": False, "reason": "input_invalid", "stage": "input"}
    else:
        answer = check(body)
    print(json.dumps(answer, ensure_ascii=False, allow_nan=False, separators=(",", ":")))
