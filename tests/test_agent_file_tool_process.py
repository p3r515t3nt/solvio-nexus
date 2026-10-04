"""Native offline file tools; all inputs/canaries/outputs are synthetic."""
from __future__ import annotations

import asyncio
import base64
from dataclasses import replace
import hashlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
from unittest.mock import patch
import zipfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
from _guard import enforce_assertions, require, require_equal
enforce_assertions()
from solvio.agent_runtime import file_tool_process as F

_BUNDLE = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/python"
_RUNTIME = None


def runtime():
    global _RUNTIME
    if _RUNTIME is None:
        # An explicit test-runtime override is independent of production's
        # interpreter. Missing packages/runtime fail this native test suite.
        root = Path(os.environ.get("SOLVIO_FILE_TOOL_TEST_PREFIX", str(_BUNDLE))).resolve()
        _RUNTIME = F.bind_runtime(python=str(root / "bin/python3"), prefix=str(root),
            site_packages=str(root / "lib/python3.12/site-packages"))
    return _RUNTIME


def invocation(root, source, **options):
    source_root = root / "source"
    source_root.mkdir(exist_ok=True)
    raw = source.encode()
    (source_root / "worker.py").write_bytes(raw)
    return F.FileToolInvocation(str(source_root.resolve()), "worker.py",
        {"worker.py": hashlib.sha256(raw).hexdigest()}, runtime(), **options)


async def run(source, payload=b"{}", **options):
    require(F.E.isolation.available(), "native sandbox required for this suite")
    with tempfile.TemporaryDirectory(prefix="file-tool-test-") as directory:
        return await F.run_file_tool(invocation(Path(directory), source, **options), payload)


async def t_json_bytes_return_from_actual_bound_runtime():
    outcome = await run("import sys\nsys.stdout.buffer.write(sys.stdin.buffer.read())", b'{"value":7}')
    require(outcome.ok, outcome.reason)
    require_equal(outcome.stdout, b'{"value":7}')
    require_equal(outcome.execution_status, "terminal")
    require(outcome.process_started)


async def t_native_spreadsheet_chart_and_pdf_bytes_are_created_and_reopened():
    source = '''import base64, io, json
import pandas as pd
import xlsxwriter
import openpyxl
from PIL import Image, ImageDraw
from reportlab.pdfgen.canvas import Canvas
from pypdf import PdfReader
frame = pd.DataFrame({'month':['Jan','Feb','Mar'], 'amount':[3,5,7]})
total = int(frame['amount'].sum())
book_data = io.BytesIO()
with xlsxwriter.Workbook(book_data, {'in_memory':True}) as book:
    sheet = book.add_worksheet('Analysis')
    sheet.write_row(0,0,['Month','Amount'])
    for row, (month, amount) in enumerate(frame.itertuples(index=False, name=None), 1):
        sheet.write_row(row,0,[month,amount])
    sheet.write(4,0,'Total')
    sheet.write_formula(4,1,'=SUM(B2:B4)', None, total)
    chart = book.add_chart({'type':'column'})
    chart.add_series({'name':'Amount','categories':'=Analysis!$A$2:$A$4',
        'values':'=Analysis!$B$2:$B$4'})
    sheet.insert_chart('D2',chart)
formula_book = openpyxl.load_workbook(io.BytesIO(book_data.getvalue()), data_only=False)
value_book = openpyxl.load_workbook(io.BytesIO(book_data.getvalue()), data_only=True)
png_data = io.BytesIO()
image = Image.new('RGB',(240,120),'white')
draw = ImageDraw.Draw(image)
for index, value in enumerate(frame['amount']):
    draw.rectangle((20+index*65,100-int(value)*10,55+index*65,100),fill='#2563eb')
image.save(png_data,format='PNG')
reopened = Image.open(io.BytesIO(png_data.getvalue()))
reopened.load()
pdf_data = io.BytesIO()
pdf = Canvas(pdf_data, pagesize=(300,200))
pdf.drawString(20,170,'Synthetic amount total: ' + str(total))
pdf.save()
pdf_text = PdfReader(io.BytesIO(pdf_data.getvalue())).pages[0].extract_text()
files = {name:base64.b64encode(data).decode() for name,data in (
    ('analysis.xlsx',book_data.getvalue()),('chart.png',png_data.getvalue()),
    ('report.pdf',pdf_data.getvalue()))}
print(json.dumps({'files':files,'total':total,'formula':formula_book['Analysis']['B5'].value,
    'cached':value_book['Analysis']['B5'].value,'size':list(reopened.size),'pdf_text':pdf_text}))
'''
    samples = []
    native_rss = F._memory_rss
    def measured(pid):
        rss = native_rss(pid)
        if rss is not None:
            samples.append(rss)
        return rss
    with patch.object(F, "_memory_rss", measured):
        outcome = await run(source)
    require(outcome.ok, outcome.reason)
    require(samples and max(samples) < F.DEFAULT_MEMORY_BYTES)
    print("Office fixture observed peak RSS bytes:", max(samples))
    body = json.loads(outcome.stdout)
    require_equal(body["total"], 15)
    require_equal(body["cached"], 15)
    require_equal(body["formula"], "=SUM(B2:B4)")
    require_equal(body["size"], [240, 120])
    require("Synthetic amount total: 15" in body["pdf_text"])
    files = {name: base64.b64decode(value, validate=True) for name, value in body["files"].items()}
    require(files["chart.png"].startswith(b"\x89PNG\r\n\x1a\n"))
    require(files["report.pdf"].startswith(b"%PDF-"))
    with zipfile.ZipFile(io.BytesIO(files["analysis.xlsx"])) as archive:
        require("xl/charts/chart1.xml" in archive.namelist())
        require(b"SUM(B2:B4)" in archive.read("xl/worksheets/sheet1.xml"))


async def t_source_change_symlink_and_escape_never_spawn():
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        call = invocation(root, "print('{}')")
        (root / "source/worker.py").write_text("print('changed')")
        changed = await F.run_file_tool(call, b"{}")
        require_equal(changed.reason, "artifact_changed")
        require(not changed.process_started)
        escaped = await F.run_file_tool(replace(call, entrypoint="../worker.py"), b"{}")
        require_equal(escaped.reason, "invalid_artifact_path")
        require(not escaped.process_started)
        (root / "source/worker.py").unlink()
        (root / "synthetic.py").write_text("print('{}')")
        (root / "source/worker.py").symlink_to(root / "synthetic.py")
        linked = await F.run_file_tool(call, b"{}")
        require(not linked.process_started)


async def t_runtime_file_and_additional_module_changes_invalidate_measured_binding():
    # Synthetic installation with a real-shaped local tree; no execution or
    # changes to the installed tool runtime are needed to prove hash coverage.
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory).resolve()
        (root / "bin").mkdir()
        binary = root / "bin/python3.12"
        binary.write_bytes(b"synthetic runtime")
        binary.chmod(0o700)
        packages = root / "lib/python3.12/site-packages"
        packages.mkdir(parents=True)
        module = packages / "module.py"
        module.write_text("number=1")
        call = invocation(root, "print('{}')")
        bound = F.bind_runtime(python=str(binary), prefix=str(root), site_packages=str(packages))
        call = replace(call, runtime=bound)
        require_equal(F._runtime_fingerprint(bound.python, bound.prefix, bound.site_packages), bound.fingerprint)
        module.write_text("number=2")
        changed = await F.run_file_tool(call, b"{}")
        require_equal(changed.reason, "runtime_changed")
        require(not changed.process_started)
        module.write_text("number=1")
        require_equal(F._runtime_fingerprint(bound.python, bound.prefix, bound.site_packages), bound.fingerprint)
        (packages / "extra.pyc").write_bytes(b"new unbound bytecode")
        changed = await F.run_file_tool(call, b"{}")
        require_equal(changed.reason, "runtime_changed")
        require(not changed.process_started)


async def t_source_snapshot_is_immutable_even_if_original_changes_after_spawn():
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        source = '''import json,pathlib
print(json.dumps({'original':pathlib.Path(__file__).read_text().startswith('import json'),
    'unbound':pathlib.Path(__file__).with_name('unbound.txt').exists()}))
'''
        call = invocation(root, source)
        (root / "source/unbound.txt").write_text("not in bound artifact")
        original = F.E.asyncio.create_subprocess_exec
        async def spawn(*args, **kwargs):
            process = await original(*args, **kwargs)
            (root / "source/worker.py").write_text("print('changed')")
            return process
        with patch.object(F.E.asyncio, "create_subprocess_exec", spawn):
            outcome = await F.run_file_tool(call, b"{}")
        require(outcome.ok, outcome.reason)
        require_equal(json.loads(outcome.stdout), {"original":True, "unbound":False})


async def t_postexecution_runtime_mismatch_withholds_actual_completed_output():
    # Real child, synthetic post-measurement mismatch. The preceding test
    # separately proves actual changed-library/added-bytecode measurement.
    with tempfile.TemporaryDirectory() as directory:
        call = invocation(Path(directory), "print('{\"answer\":7}')")
        with patch.object(F, "_runtime_fingerprint", side_effect=[call.runtime.fingerprint, "0"*64]):
            result = await F.run_file_tool(call, b"{}")
        require_equal(result.reason, "runtime_changed")
        require_equal(result.execution_status, "unknown")
        require(result.process_started)
        require_equal(result.exit_code, 0)
        require_equal(result.stdout, b"")


async def t_missing_or_malformed_runtime_and_limits_refuse_before_spawn():
    with tempfile.TemporaryDirectory() as directory:
        call = invocation(Path(directory), "print('{}')")
        for field, value in [("timeout_s", float("nan")), ("timeout_s", 61),
                ("max_output_bytes", True), ("scratch_files", 17), ("scratch_bytes", 0)]:
            result = await F.run_file_tool(replace(call, **{field:value}), b"{}")
            require(not result.process_started)
        for altered in [replace(call.runtime, fingerprint="x"),
                replace(call.runtime, prefix="/"), replace(call.runtime, python="python3")]:
            result = await F.run_file_tool(replace(call, runtime=altered), b"{}")
            require(not result.process_started)
        with patch.object(F.E.isolation, "available", return_value=False):
            result = await F.run_file_tool(call, b"{}")
        require_equal(result.reason, "sandbox_unavailable")


async def t_input_json_limits_reject_duplicates_nonfinite_and_nonobjects():
    with tempfile.TemporaryDirectory() as directory:
        call = invocation(Path(directory), "print('{}')", max_input_bytes=128)
        for payload in [b'{"x":1,"x":2}', b'{"x":NaN}', b'{"x":1e999}', b'[]', b'\xff', b'']:
            result = await F.run_file_tool(call, payload)
            require_equal(result.reason, "invalid_input_json")
            require(not result.process_started)
        result = await F.run_file_tool(call, b"x"*129)
        require_equal(result.reason, "input_limit")
        require(not result.process_started)


async def t_checker_can_receive_two_base64_packages_within_explicit_32mib_envelope():
    # 8 MiB source + 8 MiB produced-file packages require over 16 MiB after
    # base64. This tests only transport capacity, not workbook acceptance.
    payload = json.dumps({"source_b64":"eA==" * (3*1024*1024),
                          "result_b64":"eA==" * (3*1024*1024)}).encode()
    require(16*1024*1024 < len(payload) < F.MAX_INPUT_BYTES)
    result = await run("import json,sys\nx=json.load(sys.stdin)\nprint(json.dumps({'fields':len(x)}))", payload)
    require(result.ok, result.reason)
    require_equal(json.loads(result.stdout), {"fields":2})


async def t_output_invalid_json_is_terminal_failure_without_partial_bytes():
    result = await run("print('not JSON')")
    require_equal(result.reason, "invalid_output_json")
    require_equal(result.execution_status, "terminal")
    require_equal(result.stdout, b"")


async def t_stderr_and_stdout_share_budget_and_stderr_is_never_disclosed():
    result = await run("import os\nos.write(1,b'x'*40)\nos.write(2,b'y'*40)", max_output_bytes=64)
    require_equal(result.reason, "output_limit")
    require_equal(result.execution_status, "unknown")
    require_equal(result.stdout, b"")
    require_equal(result.stderr_note, "")


async def t_foreign_read_tool_write_shell_and_keychain_binary_are_denied():
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        foreign = root / "synthetic-private"
        foreign.write_text("not visible to tool")
        source = '''import json, os, pathlib, sys
foreign = json.load(sys.stdin)['foreign']
denied=[]
for action in [lambda: pathlib.Path(foreign).read_text(),
               lambda: pathlib.Path(__file__).write_text('changed'),
               lambda: os.execv('/bin/sh',['/bin/sh','-c','true']),
               lambda: os.execv('/usr/bin/security',['/usr/bin/security','help'])]:
    try: action(); denied.append(False)
    except PermissionError: denied.append(True)
print(json.dumps({'denied':denied}))
'''
        result = await F.run_file_tool(invocation(root, source), json.dumps({"foreign":str(foreign)}).encode())
        require(result.ok, result.reason)
        require_equal(json.loads(result.stdout)["denied"], [True]*4)
        require_equal(foreign.read_text(), "not visible to tool")


async def t_network_denied_even_for_a_live_synthetic_loopback_listener():
    server = await asyncio.start_server(lambda r,w:w.close(), "127.0.0.1", 0)
    try:
        port = server.sockets[0].getsockname()[1]
        _, writer = await asyncio.open_connection("127.0.0.1", port)
        writer.close()
        await writer.wait_closed()
        source = '''import json,socket,sys
port=json.load(sys.stdin)['port']
answers=[]
for address in [('127.0.0.1',port),('192.0.2.1',443)]:
    try:
        with socket.socket() as sock:
            sock.settimeout(.1)
            sock.connect(address)
        answers.append('connected')
    except PermissionError: answers.append('denied')
    except OSError: answers.append('other_error')
print(json.dumps({'answers':answers}))
'''
        result = await run(source, json.dumps({"port":port}).encode())
        require(result.ok, result.reason)
        require_equal(json.loads(result.stdout)["answers"], ["denied", "denied"])
    finally:
        server.close()
        await server.wait_closed()


async def t_environment_is_explicit_and_scratch_is_removed():
    source = "import json,os\nprint(json.dumps({'keys':sorted(os.environ),'home':os.environ['HOME'],'cwd':os.getcwd()}))"
    with patch.dict(os.environ, {"OPENAI_API_KEY":"synthetic", "PYTHONPATH":"/synthetic", "SOLVIO_VAULT_DIR":"/synthetic"}):
        result = await run(source)
    require(result.ok, result.reason)
    body = json.loads(result.stdout)
    require_equal(body["keys"], sorted(["HOME","TMPDIR","LANG","TZ","OPENBLAS_NUM_THREADS",
        "OMP_NUM_THREADS","MKL_NUM_THREADS","VECLIB_MAXIMUM_THREADS"]))
    require_equal(body["home"], body["cwd"])
    require(not Path(body["home"]).exists())


async def t_scratch_count_and_total_bytes_are_hard_bounded():
    source = '''import json,os,pathlib
sizes=[]
for path in sorted(pathlib.Path('.').glob('slot-*.tmp')):
    try:
        with path.open('wb') as file: file.write(b'x'*33); file.flush()
    except OSError: pass
    sizes.append(path.stat().st_size)
denied=[]
for action in [lambda:pathlib.Path('extra.tmp').write_bytes(b'x'),
               lambda:pathlib.Path('slot-00.tmp').unlink(),
               lambda:pathlib.Path('newdir').mkdir()]:
    try: action(); denied.append(False)
    except PermissionError: denied.append(True)
print(json.dumps({'sizes':sizes,'denied':denied}))
'''
    result = await run(source, scratch_bytes=64, scratch_files=2)
    require(result.ok, result.reason)
    require_equal(json.loads(result.stdout), {"sizes":[32,32],"denied":[True]*3})


async def t_fork_and_posix_spawn_cannot_escape_process_ownership():
    source = '''import json,os,sys
denied=[]
for action in [lambda:os.fork(),lambda:os.posix_spawn(sys.executable,
    [sys.executable,'-I','-S','-c','pass'],dict(os.environ),setsid=True)]:
    try: action(); denied.append(False)
    except PermissionError: denied.append(True)
print(json.dumps({'denied':denied}))
'''
    result = await run(source)
    require(result.ok, result.reason)
    require_equal(json.loads(result.stdout)["denied"], [True,True])


async def t_timeout_is_started_unknown_with_no_partial_json():
    result = await run("import time\ntime.sleep(30)", timeout_s=.1)
    require_equal(result.reason, "timeout")
    require_equal(result.execution_status, "unknown")
    require(result.process_started)
    require_equal(result.stdout, b"")


async def t_cancellation_during_spawn_reaps_child_and_keeps_unrelated_process():
    foreign = await asyncio.create_subprocess_exec(sys.executable,"-I","-S","-c","import time;time.sleep(30)")
    try:
        with tempfile.TemporaryDirectory() as directory:
            call = invocation(Path(directory), "import time\ntime.sleep(30)")
            original = F.E.asyncio.create_subprocess_exec
            spawned = asyncio.get_running_loop().create_future()
            release = asyncio.Event()
            scratch = []
            async def spawn(*args, **kwargs):
                process = await original(*args, **kwargs)
                scratch.append(kwargs['cwd'])
                spawned.set_result(process)
                await release.wait()
                return process
            with patch.object(F.E.asyncio,"create_subprocess_exec",spawn):
                task=asyncio.create_task(F.run_file_tool(call,b"{}"))
                process=await asyncio.wait_for(spawned,10)
                task.cancel()
                await asyncio.sleep(0)
                task.cancel()
                release.set()
                try:
                    await asyncio.wait_for(task,3)
                    require(False,"cancellation swallowed")
                except asyncio.CancelledError: pass
            require(process.returncode is not None)
            require(not Path(scratch[0]).exists())
            require(foreign.returncode is None)
    finally:
        if foreign.returncode is None: foreign.terminate()
        await foreign.wait()


async def t_memory_limit_measures_real_pages_and_reaps_only_the_owned_child():
    # At most 64 MiB of synthetic data even if the watchdog regresses. Touch
    # every page and hold it long enough for actual 25-ms RSS observations.
    source = '''import json,sys,time
sys.stdout.write('{"partial":true}'); sys.stdout.flush()
blocks=[]
for index in range(8):
 block=bytearray(8*1024*1024)
 block[::4096]=b'x'*(len(block)//4096)
 blocks.append(block)
 time.sleep(.05)
time.sleep(.5)
'''
    foreign = await asyncio.create_subprocess_exec(sys.executable, "-I", "-S", "-c", "import time;time.sleep(30)")
    owned, samples, scratch = [], [], []
    original_spawn, original_rss = F.E._spawn_owned, F._memory_rss
    async def spawn(*args, **kwargs):
        process = await original_spawn(*args, **kwargs)
        owned.append(process)
        scratch.append(kwargs["cwd"])
        return process
    def rss(pid):
        value = original_rss(pid)
        if value is not None:
            samples.append((pid, value))
        return value
    try:
        with patch.object(F.E, "_spawn_owned", spawn), patch.object(F, "_memory_rss", rss):
            outcome = await run(source, memory_bytes=F.MIN_MEMORY_BYTES, timeout_s=5)
        require_equal(outcome.reason, "memory_limit")
        require_equal(outcome.execution_status, "unknown")
        require(outcome.process_started)
        require_equal(outcome.stdout, b"")
        require_equal(outcome.stderr_note, "")
        require_equal(len(owned), 1)
        require(owned[0].returncode is not None)
        require(samples and all(pid == owned[0].pid for pid, _ in samples))
        require(max(value for _, value in samples) > F.MIN_MEMORY_BYTES)
        require(not Path(scratch[0]).exists())
        require(foreign.returncode is None)
        print("Limited fixture observed peak RSS bytes:", max(value for _, value in samples))
    finally:
        if foreign.returncode is None:
            foreign.terminate()
        await foreign.wait()


async def t_missing_memory_probe_never_releases_code_start_barrier():
    owned = []
    original_spawn = F.E._spawn_owned
    async def spawn(*args, **kwargs):
        process = await original_spawn(*args, **kwargs)
        owned.append(process)
        return process
    with patch.object(F.E, "_spawn_owned", spawn), patch.object(F, "_memory_rss", return_value=None), \
            patch.object(F.E, "_communicate", side_effect=AssertionError("code barrier released without RSS")):
        outcome = await run("print('{}')")
    require_equal(outcome.reason, "memory_measurement_unavailable")
    require_equal(outcome.execution_status, "unknown")
    require(outcome.process_started)
    require_equal(len(owned), 1)
    require(owned[0].returncode is not None)
    require_equal(outcome.stdout, b"")


async def t_probe_loss_after_start_is_held_unknown_and_owned_process_is_reaped():
    owned, count = [], 0
    original_spawn, original_rss = F.E._spawn_owned, F._memory_rss
    async def spawn(*args, **kwargs):
        process = await original_spawn(*args, **kwargs)
        owned.append(process)
        return process
    def disappearing(pid):
        nonlocal count
        count += 1
        return original_rss(pid) if count == 1 else None
    with patch.object(F.E, "_spawn_owned", spawn), patch.object(F, "_memory_rss", disappearing):
        outcome = await run("import time\ntime.sleep(30)", timeout_s=2)
    require(count >= 2)
    require_equal(outcome.reason, "memory_measurement_unavailable")
    require_equal(outcome.execution_status, "unknown")
    require_equal(outcome.stdout, b"")
    require_equal(len(owned), 1)
    require(owned[0].returncode is not None)


async def t_cancel_during_live_memory_monitor_still_reaps_after_repeated_cancel():
    monitored = asyncio.Event()
    owned, count = [], 0
    original_spawn, original_rss = F.E._spawn_owned, F._memory_rss
    async def spawn(*args, **kwargs):
        process = await original_spawn(*args, **kwargs)
        owned.append(process)
        return process
    def rss(pid):
        nonlocal count
        count += 1
        if count >= 3:
            monitored.set()
        return original_rss(pid)
    with patch.object(F.E, "_spawn_owned", spawn), patch.object(F, "_memory_rss", rss):
        task = asyncio.create_task(run("import time\ntime.sleep(30)"))
        await asyncio.wait_for(monitored.wait(), 10)
        task.cancel()
        await asyncio.sleep(0)
        task.cancel()
        try:
            await asyncio.wait_for(task, 3)
        except asyncio.CancelledError:
            pass
        else:
            require(False, "cancellation swallowed")
    require_equal(len(owned), 1)
    require(owned[0].returncode is not None)


async def t_memory_limit_is_fixed_bounded_and_never_accepted_from_untyped_values():
    with tempfile.TemporaryDirectory() as directory:
        call = invocation(Path(directory), "print('{}')")
        with patch.object(F.E, "_spawn_owned", side_effect=AssertionError("invalid limit spawned")):
            for value in (True, "536870912", F.MIN_MEMORY_BYTES - 1, F.MAX_MEMORY_BYTES + 1, float("inf")):
                outcome = await F.run_file_tool(replace(call, memory_bytes=value), b"{}")
                require_equal(outcome.reason, "invalid_memory_limit")
                require(not outcome.process_started)


async def t_cancellation_during_posthash_keeps_loop_responsive_and_child_already_reaped():
    with tempfile.TemporaryDirectory() as directory:
        call = invocation(Path(directory), "print('{}')")
        original_hash, original_spawn = F._runtime_fingerprint, F.E._spawn_owned
        loop = asyncio.get_running_loop()
        entered = asyncio.Event()
        release = threading.Event()
        finished = threading.Event()
        owned, scratch, count = [], [], 0
        def fingerprint(*args):
            nonlocal count
            count += 1
            if count == 1:
                return original_hash(*args)
            loop.call_soon_threadsafe(entered.set)
            try:
                require(release.wait(5), "posthash fixture was not released")
                return call.runtime.fingerprint
            finally:
                finished.set()
        async def spawn(*args, **kwargs):
            process = await original_spawn(*args, **kwargs)
            owned.append(process)
            scratch.append(kwargs["cwd"])
            return process
        with patch.object(F, "_runtime_fingerprint", fingerprint), patch.object(F.E, "_spawn_owned", spawn):
            task = asyncio.create_task(F.run_file_tool(call, b"{}"))
            try:
                await asyncio.wait_for(entered.wait(), 10)
                require_equal(len(owned), 1)
                require(owned[0].returncode is not None, "posthash began with a live child")
                task.cancel()
                try:
                    await asyncio.wait_for(task, 1)
                except asyncio.CancelledError:
                    pass
                else:
                    require(False, "posthash cancellation swallowed")
                require(not Path(scratch[0]).exists())
            finally:
                release.set()
                require(await asyncio.to_thread(finished.wait, 2))


if __name__ == "__main__":
    from _harness import run_module
    raise SystemExit(run_module(globals(), __name__))
