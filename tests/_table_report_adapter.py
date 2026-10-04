SOURCE = r'''"""Synthetic builder output for native table-report integration tests.

This is a test producer, not a shipped capability or a copy of the Core oracle.
It uses the libraries named by the fixed tool contract and never imports Core.
Tests may deliberately mutate these outputs before the independent readback.
"""
import base64
import csv
import io
import json
import math
import sys
import zipfile

import openpyxl
import pandas as pd
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas
import xlsxwriter


def cell(value):
    if pd.isna(value):
        return None
    if hasattr(value, "item"):
        value = value.item()
    if hasattr(value, "isoformat"):
        value = value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("nonfinite")
    if type(value) is str and len(value) > 32767:
        raise ValueError("cell_limit")
    return value


def read_tables(files):
    result, cells = [], 0
    for item in files:
        raw = base64.b64decode(item["content_b64"], validate=True)
        filename = item["name"]
        if filename.lower().endswith(".csv"):
            text = raw.decode("utf-8-sig")
            try:
                delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;").delimiter
            except csv.Error:
                if "," in text[:65536] or ";" in text[:65536]:
                    raise ValueError("csv_dialect") from None
                delimiter = ","
            header = next(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True))
            if not header or len(set(header)) != len(header) or any(not value.strip() for value in header):
                raise ValueError("headers")
            frames = [("CSV", pd.read_csv(io.BytesIO(raw), sep=delimiter, engine="python",
                                         encoding="utf-8-sig", nrows=50001))]
        elif filename.lower().endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries = archive.infolist()
                if (len(entries) > 2048 or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024
                        or any(entry.flag_bits & 1 or entry.filename.startswith("/")
                            or ".." in entry.filename.split("/")
                            or any(part in entry.filename.lower() for part in
                                ("vbaproject", "externallinks/", "embeddings/")) for entry in entries)):
                    raise ValueError("archive")
            book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            frames = []
            for sheet in book.worksheets:
                sheet.reset_dimensions()
                rows, actual_cells = [], sum(frame.size for _, frame in frames)
                for index, row in enumerate(sheet.iter_rows()):
                    if index > 50000 or len(row) > 128:
                        raise ValueError("table_limit")
                    if index:
                        actual_cells += len(row)
                        if cells + actual_cells > 500000:
                            raise ValueError("table_limit")
                    rows.append(row)
                if not rows or any(value.data_type == "f" for row in rows for value in row):
                    raise ValueError("formula_or_empty")
                frames.append((sheet.title, pd.DataFrame(
                    [[value.value for value in row] for row in rows[1:]],
                    columns=[value.value for value in rows[0]])))
            book.close()
        else:
            raise ValueError("format")
        for sheet, frame in frames:
            names = list(frame.columns)
            if (len(frame) > 50000 or not 1 <= len(names) <= 128
                    or len(set(names)) != len(names)
                    or any(type(name) is not str or not name.strip() or len(name) > 120 for name in names)):
                raise ValueError("table_shape")
            cells += len(frame) * len(names)
            if cells > 500000 or len(result) >= 8:
                raise ValueError("table_limit")
            metrics = {}
            for name in names:
                values = frame[name]
                if pd.api.types.is_numeric_dtype(values) and not pd.api.types.is_bool_dtype(values):
                    numbers = [float(number) for number in values.dropna()]
                    if numbers:
                        if not all(math.isfinite(number) for number in numbers):
                            raise ValueError("nonfinite")
                        total = math.fsum(numbers)
                        metrics[name] = dict(count=len(numbers), sum=total, min=min(numbers),
                                             max=max(numbers), mean=total / len(numbers))
            facts = {"source": filename, "sheet": sheet, "rows": len(frame), "columns": names,
                "missing": {name: int(frame[name].isna().sum()) for name in names}, "numeric": metrics}
            data = [[cell(value) for value in row] for row in frame.itertuples(index=False, name=None)]
            result.append((facts, data))
    return result


def report(tables):
    stats = [["Datei", "Tabelle", "Spalte", "Anzahl", "Summe", "Minimum", "Maximum", "Mittelwert"]]
    plotted = []
    for facts, _ in tables:
        for name, metric in facts["numeric"].items():
            stats.append([facts["source"], facts["sheet"], name] +
                [metric[key] for key in ("count", "sum", "min", "max", "mean")])
            plotted.append((name, metric["sum"]))
    workbook = io.BytesIO()
    with xlsxwriter.Workbook(workbook, {"in_memory": True,
            "strings_to_formulas": False, "strings_to_urls": False}) as book:
        header = book.add_format({"bold": True, "bg_color": "#DDEAF7", "text_wrap": True})
        numeric = book.add_format({"num_format": "0.00"})
        sheets = [("Kennzahlen", stats)] + [("Daten" + str(index), [facts["columns"]] + rows)
                    for index, (facts, rows) in enumerate(tables, 1)]
        for name, rows in sheets:
            sheet = book.add_worksheet(name)
            sheet.freeze_panes(1, 0)
            sheet.set_column(0, len(rows[0]) - 1, 20)
            for row, values in enumerate(rows):
                sheet.write_row(row, 0, values, header if row == 0 else None)
            if name == "Kennzahlen":
                sheet.set_column(3, 7, 17, numeric)
                chart = book.add_chart({"type": "column"})
                if len(stats) > 1:
                    chart.add_series({"name": "Summe", "categories": [name, 1, 2, len(stats)-1, 2],
                                      "values": [name, 1, 4, len(stats)-1, 4]})
                else:
                    chart.add_series({"name": "Zeilen", "values": "={" +
                        ",".join(str(facts["rows"]) for facts, _ in tables) + "}",
                        "values_data": [facts["rows"] for facts, _ in tables]})
                chart.set_title({"name": "Tabellenübersicht"})
                chart.set_legend({"none": True})
                sheet.insert_chart("J2", chart)

    if not plotted:
        plotted = [(facts["source"], facts["rows"]) for facts, _ in tables]
    chart = Image.new("RGB", (960, 540), "white")
    drawing = ImageDraw.Draw(chart)
    drawing.text((28, 18), "SOLVIO: Summen der numerischen Spalten / sonst Zeilenzahlen", fill="#173553")
    drawing.text((28, 38), "Einheiten nur aus Quelldaten; keine Waehrung angenommen.", fill="#384B5F")
    scale = max([abs(value) for _, value in plotted] + [1])
    shown = plotted[:12]
    for index, (label, value) in enumerate(shown):
        y = 80 + index * 34
        x = 450 + int(value / scale * 340)
        drawing.text((25, y + 3), str(label)[:50], fill="black")
        drawing.rectangle((min(x, 450), y, max(x, 450) + 1, y + 20), fill="#276FBF")
        drawing.text((800, y + 3), format(value, ".6g"), fill="black")
    image = io.BytesIO()
    chart.save(image, format="PNG")

    document = io.BytesIO()
    pdf = canvas.Canvas(document, pagesize=(595, 842), invariant=1)
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(38, 802, "SOLVIO Tabellenbericht")
    y = 770
    def line(text):
        nonlocal y
        if y < 55:
            pdf.showPage()
            y = 800
        pdf.setFont("Helvetica", 10)
        pdf.drawString(38, y, text)
        y -= 16
    for index, (facts, _) in enumerate(tables, 1):
        line("Tabelle %d: %d Zeilen" % (index, facts["rows"]))
        line("Quelle: %s / %s" % (facts["source"], facts["sheet"]))
        for column, metric in facts["numeric"].items():
            number = facts["columns"].index(column) + 1
            line("%s (Spalte %d): Summe: %s" % (column, number, format(metric["sum"], ".6g")))
        line("Fehlende Werte: " + ", ".join("%s: %s" % pair for pair in facts["missing"].items()))
        y -= 8
    line("Fehlende Werte sind aus den numerischen Kennzahlen ausgeschlossen.")
    line("Die Uebersicht erlaubt keine Schlussfolgerung zu Korrelationen oder Ursachen.")
    pdf.save()
    payloads = [("Analyse.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", workbook.getvalue()),
                ("Diagramm.png", "image/png", image.getvalue()),
                ("Bericht.pdf", "application/pdf", document.getvalue())]
    if sum(len(raw) for _, _, raw in payloads) > 8 * 1024 * 1024:
        raise ValueError("output_limit")
    return {"version": 1, "summary": "Die Tabellenübersicht mit Kennzahlen, Arbeitsmappe, Diagramm und Bericht ist erstellt.",
        "tables": [facts for facts, _ in tables], "files": [
            {"name": name, "mime_type": mime, "content_b64": base64.b64encode(raw).decode("ascii")}
            for name, mime, raw in payloads]}


if __name__ == "__main__":
    body = json.load(sys.stdin)
    if body["version"] != 1 or body["operation"] != "table_report":
        raise ValueError("request")
    print(json.dumps(report(read_tables(body["files"])), ensure_ascii=False,
                     allow_nan=False, separators=(",", ":")))
'''

if __name__ == "__main__":
    exec(compile(SOURCE, __file__, "exec"))
